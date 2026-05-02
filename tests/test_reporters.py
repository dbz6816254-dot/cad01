"""리포터 테스트 (Excel 파일 산출 확인)."""

from __future__ import annotations

from openpyxl import load_workbook

from rebar_audit.fixtures import mismatched_invoice_items, reference_items
from rebar_audit.matcher import compare_by_diameter, compare_by_mark
from rebar_audit.models import RebarSet
from rebar_audit.reporters.excel import write_report


def _set(items, origin="schedule"):
    return RebarSet(source_path="test", source_kind="excel", origin=origin, items=items)


def test_excel_report_creates_workbook(tmp_path):
    a = _set(reference_items())
    b = _set(mismatched_invoice_items(), origin="invoice")
    s1 = compare_by_diameter(a, b)
    s2 = compare_by_mark(a, b)
    out = tmp_path / "report.xlsx"
    write_report(out, a, b, s1, s2)

    wb = load_workbook(out)
    assert "요약" in wb.sheetnames
    assert "Stage1_직경별" in wb.sheetnames
    assert "Stage2_부호별" in wb.sheetnames
    # Stage1 시트에 헤더 + 데이터 행이 있는지
    ws = wb["Stage1_직경별"]
    assert ws.max_row >= 2  # header + at least one data row
