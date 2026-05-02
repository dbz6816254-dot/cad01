"""Excel 리포터 (openpyxl). 비교 결과를 다중 시트 워크북으로 출력."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rebar_audit.matcher import DiameterDiff, MarkDiff
from rebar_audit.models import RebarSet

_FILL_BY_STATUS = {
    "ok": PatternFill("solid", fgColor="C6EFCE"),
    "warning": PatternFill("solid", fgColor="FFEB9C"),
    "error": PatternFill("solid", fgColor="FFC7CE"),
    "missing_a": PatternFill("solid", fgColor="FFC7CE"),
    "missing_b": PatternFill("solid", fgColor="FFC7CE"),
}
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_header(ws, headers: list[str]) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            v = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(v))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 40)


def write_report(
    output: str | Path,
    a: RebarSet,
    b: RebarSet,
    stage1: list[DiameterDiff],
    stage2: list[MarkDiff],
) -> Path:
    wb = Workbook()
    _write_summary(wb.active, a, b, stage1, stage2)
    _write_stage1(wb.create_sheet("Stage1_직경별"), stage1)
    _write_stage2(wb.create_sheet("Stage2_부호별"), stage2)
    _write_raw(wb.create_sheet("A_원본"), a)
    _write_raw(wb.create_sheet("B_원본"), b)
    out = Path(output)
    wb.save(out)
    return out


def _write_summary(ws, a, b, stage1, stage2) -> None:
    ws.title = "요약"
    _write_header(ws, ["항목", "A", "B"])
    ws.append(["출처", a.origin, b.origin])
    ws.append(["종류", a.source_kind, b.source_kind])
    ws.append(["파일", a.source_path, b.source_path])
    ws.append(["항목 수", len(a.items), len(b.items)])
    ws.append(["총 개수", a.total_count(), b.total_count()])
    ws.append(["총 길이(m)", round(a.total_length_m(), 2), round(b.total_length_m(), 2)])
    ws.append(["총 중량(kg)", round(a.total_weight_kg(), 2), round(b.total_weight_kg(), 2)])
    ws.append([])
    ws.append(["판정 항목", "건수"])
    ws.append(["Stage1 일치", sum(1 for d in stage1 if d.status == "ok")])
    ws.append(["Stage1 경고", sum(1 for d in stage1 if d.status == "warning")])
    ws.append(
        [
            "Stage1 불일치",
            sum(1 for d in stage1 if d.status in ("error", "missing_a", "missing_b")),
        ]
    )
    ws.append(["Stage2 일치", sum(1 for d in stage2 if d.status == "ok")])
    ws.append(["Stage2 불일치", sum(1 for d in stage2 if d.status != "ok")])
    _autosize(ws)


def _write_stage1(ws, diffs: list[DiameterDiff]) -> None:
    headers = [
        "직경",
        "A 개수",
        "B 개수",
        "A 길이(m)",
        "B 길이(m)",
        "A 중량(kg)",
        "B 중량(kg)",
        "Δ중량(kg)",
        "Δ%",
        "상태",
    ]
    _write_header(ws, headers)
    for d in diffs:
        row = [
            f"D{d.diameter}",
            d.a_count,
            d.b_count,
            d.a_length_m,
            d.b_length_m,
            d.a_weight_kg,
            d.b_weight_kg,
            d.weight_delta_kg,
            d.weight_delta_pct,
            d.status,
        ]
        ws.append(row)
        fill = _FILL_BY_STATUS.get(d.status)
        if fill:
            for col in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col).fill = fill
    _autosize(ws)


def _write_stage2(ws, diffs: list[MarkDiff]) -> None:
    headers = [
        "부호",
        "A 직경",
        "A 길이",
        "A 개수",
        "B 직경",
        "B 길이",
        "B 개수",
        "상태",
        "비고",
    ]
    _write_header(ws, headers)
    for d in diffs:
        row = [
            d.mark,
            d.a.diameter if d.a else None,
            d.a.length_mm if d.a else None,
            d.a.count if d.a else None,
            d.b.diameter if d.b else None,
            d.b.length_mm if d.b else None,
            d.b.count if d.b else None,
            d.status,
            "; ".join(d.notes),
        ]
        ws.append(row)
        fill = _FILL_BY_STATUS.get(d.status)
        if fill:
            for col in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col).fill = fill
    _autosize(ws)


def _write_raw(ws, rs: RebarSet) -> None:
    headers = ["부호", "직경", "길이(mm)", "개수", "총길이(m)", "총중량(kg)", "형상", "재질"]
    _write_header(ws, headers)
    for item in rs.items:
        ws.append(
            [
                item.mark,
                item.diameter,
                item.length_mm,
                item.count,
                round(item.total_length_m, 3),
                round(item.total_weight_kg, 3),
                item.bend_type,
                item.grade,
            ]
        )
    _autosize(ws)
